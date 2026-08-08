from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font

from database.db import get_session
from database.models import Facture, Paiement


def _invoice_rows(session):
    factures = session.query(Facture).order_by(Facture.date_facture.desc()).all()
    rows = []
    for facture in factures:
        paiements = session.query(Paiement).filter_by(facture_id=facture.id).order_by(Paiement.date_paiement.asc()).all()
        rows.append({
            "ID": facture.id,
            "Numéro": facture.numero,
            "Fournisseur": facture.fournisseur.nom if facture.fournisseur else "",
            "Date facture": facture.date_facture,
            "Date échéance": facture.date_echeance,
            "Montant": facture.montant,
            "Montant payé": facture.montant_paye,
            "Reste": facture.reste,
            "Statut": getattr(facture.statut, "value", facture.statut),
            "Nombre paiements": len(paiements),
            "Dernier paiement": paiements[-1].date_paiement if paiements else None,
            "Commentaire": facture.commentaire or "",
        })
    return rows


def exporter_factures_excel(filepath):
    path = Path(filepath)
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")

    session = get_session()
    try:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Factures"

        headers = [
            "ID", "Numéro", "Fournisseur", "Date facture", "Date échéance",
            "Montant", "Montant payé", "Reste", "Statut", "Nombre paiements",
            "Dernier paiement", "Commentaire"
        ]
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)

        for row in _invoice_rows(session):
            sheet.append([row[h] for h in headers])

        for column in ("D", "E", "K"):
            for cell in sheet[column][1:]:
                if cell.value is not None:
                    cell.number_format = "yyyy-mm-dd"
        for column in ("F", "G", "H"):
            for cell in sheet[column][1:]:
                cell.number_format = '#,##0.00'

        widths = {"A": 8, "B": 20, "C": 30, "D": 14, "E": 16, "F": 16, "G": 16, "H": 16, "I": 24, "J": 18, "K": 18, "L": 40}
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        workbook.save(path)
        return path
    finally:
        session.close()


def creer_modele_factures_excel(filepath):
    """Crée un modèle d'import vierge avec une feuille d'instructions."""
    path = Path(filepath)
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Factures"
    headers = [
        "Numero", "Fournisseur", "Date facture", "Date echeance", "Montant",
        "Montant paye", "Date paiement", "Mode paiement", "Reference paiement", "Commentaire"
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for column in range(1, len(headers) + 1):
        sheet.column_dimensions[chr(64 + column)].width = 22
    sheet.freeze_panes = "A2"

    instructions = workbook.create_sheet("Instructions")
    instructions.append(["Colonne", "Obligatoire", "Description"])
    for cell in instructions[1]:
        cell.font = Font(bold=True)
    rows = [
        ("Numero", "Oui", "Numéro unique de la facture"),
        ("Fournisseur", "Oui", "Nom du fournisseur"),
        ("Date facture", "Oui", "AAAA-MM-JJ ou JJ/MM/AAAA"),
        ("Date echeance", "Non", "Si vide, la date de facture sera utilisée"),
        ("Montant", "Oui", "Montant total de la facture"),
        ("Montant paye", "Non", "Montant déjà payé, de 0 au montant total"),
        ("Date paiement", "Non", "Date du paiement historique"),
        ("Mode paiement", "Non", "Espèces, Chèque, Virement, Carte, etc."),
        ("Reference paiement", "Non", "Référence du règlement"),
        ("Commentaire", "Non", "Note libre"),
    ]
    for row in rows:
        instructions.append(row)
    instructions.column_dimensions["A"].width = 24
    instructions.column_dimensions["B"].width = 16
    instructions.column_dimensions["C"].width = 65
    instructions.freeze_panes = "A2"

    workbook.save(path)
    return path
