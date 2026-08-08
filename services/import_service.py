from datetime import date, datetime, timedelta
from pathlib import Path
import re
import unicodedata

from openpyxl import load_workbook

from database.db import get_session
from database.models import Facture, Fournisseur, Paiement, Historique, JournalAction, StatutFacture

ALIASES = {
    "numero": {"numero", "numero de facture", "n° facture", "no facture", "n facture", "facture", "reference facture", "ref facture"},
    "fournisseur": {"fournisseur", "nom fournisseur", "supplier", "vendor"},
    "date_facture": {"date facture", "date de facturation", "date_facturation", "date", "date emission", "date d emission"},
    "date_echeance": {"date echeance", "echeance", "date paiement limite", "date limite"},
    "montant": {"montant", "somme", "total", "montant total", "ttc", "total ttc"},
    "montant_paye": {"montant paye", "paye", "deja paye", "total paye", "regle"},
    "date_paiement": {"date paiement", "date du paiement"},
    "mode_paiement": {"mode paiement", "mode", "moyen paiement"},
    "reference_paiement": {"reference paiement", "ref paiement", "reference reglement"},
    "commentaire": {"commentaire", "note", "notes", "observation", "observations"},
    "cheque": {"cheque", "cheque n", "cheque numero", "cheque no"},
}


def _normalize(value):
    text = "" if value is None else str(value).strip().lower()
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", text)


def _parse_date(value, field):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"{field}: date invalide « {value} »")


def _parse_amount(value, field):
    if value in (None, ""):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("DA", "").replace("da", "").replace(" ", "")
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".") if text.rfind(",") > text.rfind(".") else text.replace(",", "")
    else:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError as exc:
        raise ValueError(f"{field}: montant invalide « {value} »") from exc


def _header_map(headers):
    normalized = {_normalize(h): i for i, h in enumerate(headers) if h not in (None, "")}
    result = {}
    for field, aliases in ALIASES.items():
        for alias in aliases:
            key = _normalize(alias)
            if key in normalized:
                result[field] = normalized[key]
                break
    return result


def _cell(row, mapping, field, default=None):
    index = mapping.get(field)
    return default if index is None or index >= len(row) else row[index]


def _infer_supplier(filepath):
    stem = _normalize(Path(filepath).stem)
    if "chark" in stem:
        return "Chark Pharma"
    if "ciphaco" in stem:
        return "CIPHACO"
    if "m sante" in stem or "m_sante" in stem or "msante" in stem:
        return "M-SANTE"
    return None


def _find_header(rows):
    for index, row in enumerate(rows[:12]):
        mapping = _header_map(row)
        if "numero" in mapping and "date_facture" in mapping and "montant" in mapping:
            return index, mapping
    return None, None


def _read_workbook(filepath):
    path = Path(filepath)
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Le fichier doit être au format Excel .xlsx")
    workbook = load_workbook(path, read_only=True, data_only=True)
    records = []
    try:
        inferred_supplier = _infer_supplier(filepath)
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            header_index, mapping = _find_header(rows)
            if mapping is None:
                continue
            for excel_row, row in enumerate(rows[header_index + 1:], start=header_index + 2):
                if not any(value not in (None, "") for value in row):
                    continue
                records.append({
                    "sheet": sheet.title,
                    "excel_row": excel_row,
                    "mapping": mapping,
                    "row": row,
                    "supplier_hint": inferred_supplier,
                })
        if not records:
            raise ValueError("Aucune feuille Excel contenant les colonnes facture/date/montant n'a été trouvée.")
        return records
    finally:
        workbook.close()


def _record_to_data(record):
    row, mapping = record["row"], record["mapping"]
    numero = str(_cell(row, mapping, "numero", "")).strip()
    supplier_name = str(_cell(row, mapping, "fournisseur", record["supplier_hint"] or "")).strip()
    if not numero:
        raise ValueError("Numéro de facture vide")
    if not supplier_name:
        raise ValueError("Fournisseur absent et impossible à déduire du nom du fichier")
    date_facture = _parse_date(_cell(row, mapping, "date_facture"), "Date facture")
    if date_facture is None:
        raise ValueError("Date facture vide")
    raw_amount = _parse_amount(_cell(row, mapping, "montant"), "Montant")
    if raw_amount == 0:
        raise ValueError("Montant nul")
    type_piece = "Avoir" if raw_amount < 0 or _normalize(numero).startswith(("avoir", "av/", "avoir/")) else "Facture"
    montant = abs(raw_amount)
    if type_piece == "Avoir":
        montant = -montant
    montant_paye = _parse_amount(_cell(row, mapping, "montant_paye", 0), "Montant payé")
    if montant_paye < 0:
        raise ValueError("Le montant payé ne peut pas être négatif")
    if type_piece == "Avoir":
        montant_paye = 0.0
    elif montant_paye > montant:
        raise ValueError("Le montant payé dépasse le montant de la facture")
    date_echeance = _parse_date(_cell(row, mapping, "date_echeance", None), "Date échéance") or date_facture
    cheque = str(_cell(row, mapping, "cheque", "") or "").strip()
    commentaire = str(_cell(row, mapping, "commentaire", "") or "").strip()
    if cheque:
        commentaire = f"{commentaire} | Chèque: {cheque}" if commentaire else f"Chèque: {cheque}"
    return {
        "sheet": record["sheet"], "excel_row": record["excel_row"], "numero": numero,
        "fournisseur": supplier_name, "date_facture": date_facture, "date_echeance": date_echeance,
        "montant": montant, "montant_paye": montant_paye, "type_piece": type_piece,
        "commentaire": commentaire, "cheque": cheque,
        "date_paiement": _parse_date(_cell(row, mapping, "date_paiement", None), "Date paiement") if _cell(row, mapping, "date_paiement", None) not in (None, "") else None,
        "mode_paiement": str(_cell(row, mapping, "mode_paiement", "") or "").strip(),
        "reference_paiement": str(_cell(row, mapping, "reference_paiement", "") or "").strip(),
    }


def analyser_factures_excel(filepath):
    """Analyse un Excel sans écrire dans la base. Retourne un aperçu exploitable par l'UI."""
    records = _read_workbook(filepath)
    session = get_session()
    try:
        existing_numbers = {n for (n,) in session.query(Facture.numero).all()}
        seen = set()
        rows, errors = [], []
        for record in records:
            try:
                data = _record_to_data(record)
                key = data["numero"]
                if key in existing_numbers:
                    data["status"] = "Doublon base"
                    data["action"] = "Ignoré"
                elif key in seen:
                    data["status"] = "Doublon Excel"
                    data["action"] = "Ignoré"
                else:
                    data["status"] = "OK"
                    data["action"] = "Import"
                    seen.add(key)
                rows.append(data)
            except Exception as exc:
                errors.append(f"Feuille {record['sheet']} — ligne {record['excel_row']} : {exc}")
        imports = [r for r in rows if r["action"] == "Import"]
        return {
            "filename": Path(filepath).name,
            "rows": rows,
            "errors": errors,
            "total": len(records),
            "valid": len(rows),
            "to_import": len(imports),
            "duplicates": len(rows) - len(imports),
            "avoirs": sum(1 for r in imports if r["type_piece"] == "Avoir"),
            "factures": sum(1 for r in imports if r["type_piece"] == "Facture"),
            "total_factures": sum(r["montant"] for r in imports if r["type_piece"] == "Facture"),
            "total_avoirs": sum(abs(r["montant"]) for r in imports if r["type_piece"] == "Avoir"),
            "suppliers": sorted({r["fournisseur"] for r in imports}),
        }
    finally:
        session.close()


def importer_factures_excel(filepath, utilisateur=None, creer_fournisseurs=True):
    """Importe les lignes validées par analyser_factures_excel dans une transaction."""
    preview = analyser_factures_excel(filepath)
    session = get_session()
    imported = skipped = suppliers_created = payments_created = 0
    errors = list(preview["errors"])
    username = getattr(utilisateur, "username", "Invité")
    user_id = getattr(utilisateur, "id", None)
    try:
        supplier_cache = {_normalize(f.nom): f for f in session.query(Fournisseur).all()}
        existing_numbers = {n for (n,) in session.query(Facture.numero).all()}
        for data in preview["rows"]:
            if data["action"] != "Import" or data["numero"] in existing_numbers:
                skipped += 1
                continue
            try:
                key = _normalize(data["fournisseur"])
                fournisseur = supplier_cache.get(key)
                if fournisseur is None:
                    if not creer_fournisseurs:
                        raise ValueError(f"Fournisseur introuvable : {data['fournisseur']}")
                    fournisseur = Fournisseur(nom=data["fournisseur"])
                    session.add(fournisseur)
                    session.flush()
                    supplier_cache[key] = fournisseur
                    suppliers_created += 1
                montant = data["montant"]
                montant_paye = data["montant_paye"]
                reste = montant if data["type_piece"] == "Avoir" else max(0.0, montant - montant_paye)
                statut = StatutFacture.PAYEE if data["type_piece"] == "Avoir" or reste == 0 else StatutFacture.PARTIELLE if montant_paye > 0 else StatutFacture.IMPAYEE
                facture = Facture(numero=data["numero"], fournisseur_id=fournisseur.id, date_facture=data["date_facture"], date_echeance=data["date_echeance"], montant=montant, montant_paye=montant_paye, reste=reste, statut=statut, type_piece=data["type_piece"], commentaire=data["commentaire"], created_by=user_id)
                session.add(facture)
                session.flush()
                session.add(Historique(facture_id=facture.id, action="Import Excel", details=f"{data['type_piece']} {data['numero']} importé depuis {preview['filename']} / {data['sheet']}", utilisateur=username))
                # Un champ « chèque » sans montant ne devient pas automatiquement un paiement.
                if montant_paye > 0:
                    date_paiement = data["date_paiement"] or data["date_facture"]
                    session.add(Paiement(facture_id=facture.id, montant=montant_paye, date_paiement=date_paiement, mode_paiement=data["mode_paiement"] or ("Chèque" if data["cheque"] else "Historique"), reference=data["reference_paiement"] or data["cheque"] or None, utilisateur_id=user_id))
                    payments_created += 1
                existing_numbers.add(data["numero"])
                imported += 1
            except Exception as exc:
                errors.append(f"{data['sheet']} — ligne {data['excel_row']} — {data['numero']} : {exc}")
        if imported == 0 and errors:
            raise ValueError("Aucune ligne importable.\n\n" + "\n".join(errors[:15]))
        session.add(JournalAction(action="Import Excel", description=f"{imported} pièce(s) importée(s) depuis {preview['filename']}; {skipped} doublon(s); {suppliers_created} fournisseur(s) créé(s).", utilisateur=username))
        session.commit()
        return {"imported": imported, "skipped": skipped, "errors": errors, "suppliers_created": suppliers_created, "payments_created": payments_created, "filename": preview["filename"], "factures": preview["factures"], "avoirs": preview["avoirs"]}
    except Exception:
        session.rollback()
        raise
    finally:
        session.close()
